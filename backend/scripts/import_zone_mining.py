"""矿业勘探·开采专区(mining)商品导入:6 张配货表 → 专区 + 平台商品(ZONE_ONLY)。

输入(目录由 --src 指定,默认仓库外的 矿业专区资料/):
- 01_浮选厂*.xlsx / 02_井下开采*.xlsx / 03_矿山化验室*.xlsx / 04_金矿CIP_CIL*.xlsx / 05_露天开采*.xlsx / 600型与800型*.xlsx
- data/zone_import/mining/spu_merge.csv   原品名 → SPU 品名 + 规格(规格留空 = 纯别名)
- data/zone_import/mining/rules.csv       (file_key, sheet, section, keyword) → (专区大类, 平台品类)

产出:
- zone(code=mining)+ 11 个 kind=CATEGORY 大类 + 6 个 kind=SCENE 场景
- 平台一级 33/34/35 及其叶子(整棵 is_active=false)
- 每个 SPU:products(ZONE_ONLY/ACTIVE)+ SKU(多规格走变体)+ SPU 级属性 + 1 行 CATEGORY zone_product + N 行 SCENE zone_product
- --dry-run 写审核 CSV:data/zone_import/mining/dry_run_audit.csv + 丢弃行 dry_run_dropped.csv(不提交到 git)

幂等:zone/zone_category 按 code upsert;product 按 spu_code;zone_product 按三元组;已存在商品只更新可变字段、不动 SKU。
规格变更(spu_merge.csv 加/删规格)不会自动增删 SKU:脚本对比库内 SKU 与期望集合,不一致就拒绝落库,需人工处理。

用法:
    cd backend
    python scripts/import_zone_mining.py --src "/path/矿业专区资料" --dry-run
    python scripts/import_zone_mining.py --src "/path/矿业专区资料" --commit
--commit 且 DB 不在本机(staging/生产容器)须 ALLOW_ZONE_IMPORT_PROD=1。
"""
from __future__ import annotations

import argparse
import asyncio
import collections
import csv
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

_BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_BACKEND_ROOT))
_REPO_ROOT = _BACKEND_ROOT.parent

import openpyxl  # noqa: E402

from app.services.product_code import normalize_source_key, platform_spu_code, platform_sku_code  # noqa: E402
from scripts._zone_import_common import guard_commit  # noqa: E402

ZONE_CODE = "mining"
ZONE_NAME = ("矿业勘探·开采", "Mining Exploration & Extraction", "Uchunguzi na Uchimbaji wa Madini")
SOURCE_KEY = "ZMINE"   # 统一 code 规则里的来源标识,只进 hash
SOURCE = "IMPORT"
BATCH_ID = "zone_mining_v1"

DATA_DIR = _REPO_ROOT / "data" / "zone_import" / "mining"
MERGE_CSV = DATA_DIR / "spu_merge.csv"
RULES_CSV = DATA_DIR / "rules.csv"
AUDIT_CSV = DATA_DIR / "dry_run_audit.csv"
DROPPED_CSV = DATA_DIR / "dry_run_dropped.csv"
MAX_NAME_LEN = 60

# 不导入的 sheet(汇总/依据/重复)
SKIP_SHEETS = {
    "规模与汇总", "原表处理对照", "依据与说明", "口径与来源",
    "预算汇总", "原始预算依据", "新报价依据", "800型10台预算",
}

# 文件前缀 → 场景 code
SCENE_BY_FILE_KEY = {"01": "S01", "02": "S02", "03": "S05", "04": "S04", "05": "S03", "600": "S06"}

# (code, zh, en, sw)
ZONE_CATEGORIES = [
    ("C01", "破碎与磨矿", "Crushing & Grinding", "Kusaga na Kuponda"),
    ("C02", "选矿与浸出设备", "Beneficiation & Leaching Equipment", "Vifaa vya Uchenjuaji"),
    ("C03", "钻探与凿岩", "Drilling & Rock Drilling", "Uchimbaji Visima"),
    ("C04", "矿用药剂与工艺耗材", "Mining Reagents & Process Consumables", "Kemikali za Uchimbaji"),
    ("C05", "泵、管路与输送", "Pumps, Piping & Conveying", "Pampu, Mabomba na Usafirishaji"),
    ("C06", "电气、照明与仪表", "Electrical, Lighting & Instruments", "Umeme, Taa na Vipimo"),
    ("C07", "工程机械与车辆易损件", "Machinery & Vehicle Wear Parts", "Vipuri vya Mitambo na Magari"),
    ("C08", "油品与润滑", "Oils & Lubricants", "Mafuta na Vilainishi"),
    ("C09", "维修工具", "Maintenance Tools", "Zana za Matengenezo"),
    ("C10", "劳保与现场耗材", "PPE & Site Consumables", "Vifaa vya Usalama na Matumizi"),
    ("C11", "化验室仪器与耗材", "Laboratory Instruments & Consumables", "Vifaa vya Maabara"),
]
ZONE_SCENES = [
    ("S01", "浮选厂", "Flotation Plant", "Kiwanda cha Flotation"),
    ("S02", "井下开采", "Underground Mining", "Uchimbaji wa Chini ya Ardhi"),
    ("S03", "露天开采", "Open-pit Mining", "Uchimbaji wa Wazi"),
    ("S04", "金矿 CIP·CIL", "Gold CIP/CIL", "Dhahabu CIP/CIL"),
    ("S05", "矿山化验室", "Mine Laboratory", "Maabara ya Mgodi"),
    ("S06", "勘探钻机", "Exploration Drill Rig", "Mashine ya Uchimbaji wa Utafiti"),
]
# (code, zh, en, sw, parent_code)  整棵 is_active=false
PLATFORM_CATEGORIES = [
    ("33", "矿山与钻探机械", "Mining & Drilling Machinery", "Mitambo ya Uchimbaji", None),
    ("33.001", "破碎与磨矿", "Crushing & Grinding", "Kusaga na Kuponda", "33"),
    ("33.002", "选矿与浸出", "Beneficiation & Leaching", "Uchenjuaji", "33"),
    ("33.003", "钻探与凿岩", "Drilling & Rock Drilling", "Uchimbaji Visima", "33"),
    ("33.004", "矿用泵与输送", "Mining Pumps & Conveying", "Pampu na Usafirishaji", "33"),
    ("34", "工业化学品与药剂", "Industrial Chemicals & Reagents", "Kemikali za Viwandani", None),
    ("34.001", "选矿药剂", "Flotation Reagents", "Kemikali za Flotation", "34"),
    ("34.002", "浸出与回收药剂", "Leaching & Recovery Reagents", "Kemikali za Uchenjuaji", "34"),
    ("35", "实验室仪器与耗材", "Laboratory Instruments & Consumables", "Vifaa vya Maabara", None),
    ("35.001", "分析仪器", "Analytical Instruments", "Vifaa vya Uchambuzi", "35"),
    ("35.002", "制样设备", "Sample Preparation Equipment", "Vifaa vya Kuandaa Sampuli", "35"),
    ("35.003", "器皿与耗材", "Labware & Consumables", "Vyombo na Matumizi ya Maabara", "35"),
    ("35.004", "试剂", "Reagents", "Vitendanishi", "35"),
]


# ── 纯函数 ────────────────────────────────────────────────

def hard(s: str) -> str:
    """硬归一化(与央企脚本共用 product_code.normalize_source_key)。join key 与稳定 code 派生都用它。"""
    return normalize_source_key(s)


def mining_spu_code(name: str) -> str:
    return platform_spu_code(SOURCE_KEY, hard(name))


def mining_sku_code(spu_code: str, spec: str | None) -> str:
    return platform_sku_code(SOURCE_KEY, f"{spu_code}:{hard(spec) if spec else 'D'}")


@dataclass
class SourceRow:
    file_key: str      # 01 / 02 / 03 / 04 / 05 / 600
    sheet: str
    row_id: str        # 表内编号列
    section: str       # 环节/类别(600 表无此列 → "")
    name: str
    spec_text: str     # 参考型号/规格(采购说明文本)
    unit: str
    generic: str       # 是否通用


def _file_key(path: Path) -> str | None:
    n = path.name
    if n[:2].isdigit() and n[2] == "_":
        return n[:2]
    if n.startswith("600"):
        return "600"
    return None


def _cell(v) -> str:
    return str(v).strip() if v not in (None, "") else ""


_NOISE_RE = re.compile(r"^(待核|待定|合计|小计)|(方案|服务|报表)$")


# 源表里表示"没填"的占位值,当空处理,不能展示给买家。导入后按字段 distinct 扫描发现新的就加进来。
_SPEC_PLACEHOLDERS = frozenset({"未列", "原PDF未列规格"})   # 规格/说明列 → 商品描述
_GENERIC_PLACEHOLDERS = frozenset({"待确认"})             # 是否通用列 → SPU 属性


def clean_spec_text(v: str) -> str:
    return "" if v.strip() in _SPEC_PLACEHOLDERS else v


def clean_generic(v: str) -> str:
    return "" if v.strip() in _GENERIC_PLACEHOLDERS else v


def is_noise_name(name: str) -> bool:
    """表里混着的占位/服务行(「待核数量」「5台挖机方案」「车辆拖救服务」),不是商品。"""
    return bool(_NOISE_RE.search(name))


@dataclass
class DroppedRow:
    """被过滤掉的源行。丢弃必须留痕(计数 + CSV),否则真商品被误判成噪声没人知道。"""
    file_key: str
    sheet: str
    row_id: str
    name: str
    reason: str   # noise / too_long


def drop_reason(name: str) -> str | None:
    if len(name) > MAX_NAME_LEN:
        return "too_long"
    if is_noise_name(name):
        return "noise"
    return None


def load_source_rows(src_dir: Path, dropped: list[DroppedRow] | None = None) -> list[SourceRow]:
    """读目录下 6 个 xlsx 的商品行。表头 = 含「编号」单元格的那一行;列按表头关键字定位。

    dropped 传 list 时,被噪声/超长规则过滤的行会追加进去(编号或品名为空的空行不算)。
    """
    rows: list[SourceRow] = []
    for path in sorted(src_dir.glob("*.xlsx")):
        key = _file_key(path)
        if key is None:
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            if ws.title in SKIP_SHEETS:
                continue
            header: list[str] | None = None
            for r in ws.iter_rows(min_row=1, max_row=12, values_only=True):
                if any(_cell(c) == "编号" for c in r):
                    header = [_cell(c) for c in r]
                    break
            if header is None:
                continue

            def col(*keys: str) -> int | None:
                for k in keys:
                    for i, h in enumerate(header):
                        if h and k in h:
                            return i
                return None

            i_id, i_name = col("编号"), col("易耗品", "名称")
            i_sec, i_spec, i_unit, i_gen = col("类别"), col("型号", "规格"), col("单位"), col("是否通用")
            if i_name is None:
                continue
            started = False
            for r in ws.iter_rows(values_only=True):
                if not started:
                    started = any(_cell(c) == "编号" for c in r)
                    continue

                def at(i: int | None) -> str:
                    # read_only 模式下行长度可能短于表头,越界一律当空
                    return _cell(r[i]) if i is not None and i < len(r) else ""

                rid, name = at(i_id), at(i_name)
                if not rid or not name:
                    continue
                reason = drop_reason(name)
                if reason is not None:
                    if dropped is not None:
                        dropped.append(DroppedRow(key, ws.title, rid, name, reason))
                    continue
                rows.append(SourceRow(
                    file_key=key, sheet=ws.title, row_id=rid,
                    section=at(i_sec), name=name, spec_text=clean_spec_text(at(i_spec)),
                    unit=at(i_unit), generic=clean_generic(at(i_gen)),
                ))
        wb.close()
    return rows


def load_merge_map(path: Path) -> dict[str, tuple[str, str]]:
    out: dict[str, tuple[str, str]] = {}
    with path.open(encoding="utf-8") as f:
        for row in csv.DictReader(f):
            out[hard(row["原品名"])] = (row["SPU品名"].strip(), row["规格"].strip())
    return out


@dataclass(frozen=True)
class Rule:
    file_key: str
    sheet: str
    section: str
    keyword: str
    zone_cat: str
    category_code: str


def load_rules(path: Path) -> list[Rule]:
    rules: list[Rule] = []
    with path.open(encoding="utf-8") as f:
        for row in csv.DictReader(f):
            rules.append(Rule(
                row["file_key"].strip(), row["sheet"].strip(), row["section"].strip(),
                row["keyword"].strip(), row["zone_cat"].strip(), row["category_code"].strip(),
            ))
    return rules


def _rule_scope_matches(rule: Rule, row: SourceRow) -> bool:
    return (
        rule.file_key in ("*", row.file_key)
        and rule.sheet in ("*", row.sheet)
        and rule.section in ("*", "", row.section)
    )


def match_rule(row: SourceRow, rules: list[Rule]) -> Rule | None:
    """优先级:文件专属关键词规则 > 通配(*)关键词规则 > 无关键词规则;同级按文件顺序第一条命中。"""
    keyword_rules = [r for r in rules if r.keyword]
    for specific in (True, False):
        for rule in keyword_rules:
            if (rule.file_key != "*") is specific and rule.keyword in row.name and _rule_scope_matches(rule, row):
                return rule
    for rule in rules:
        if not rule.keyword and _rule_scope_matches(rule, row):
            return rule
    return None




@dataclass
class SpuRow:
    name: str
    spu_code: str
    # 源表「参考型号/规格」列 = 写给采购员的说明(按…核对/预算示例/待核),只进运营可见的采购备注,不当商品描述
    purchase_note: str
    specs: list[str] = field(default_factory=list)
    unit: str = ""
    section: str = ""
    generic: str = ""
    scenes: list[str] = field(default_factory=list)
    zone_cat: str | None = None
    category_code: str | None = None
    source_ids: list[str] = field(default_factory=list)
    conflicts: list[str] = field(default_factory=list)   # 大类/品类冲突:阻断落库
    attr_diffs: list[str] = field(default_factory=list)  # 环节/是否通用跨行不一致:只记审核,不阻断
    rule_desc: str = ""   # 命中规则 file/sheet/section/keyword,审核 CSV 用


def build_spus(rows: list[SourceRow], merge_map: dict[str, tuple[str, str]], rules: list[Rule]) -> list[SpuRow]:
    """源行 → SPU:按合并表改名/取规格,按归一化品名聚合;场景取并集;大类取第一次命中,不一致记 conflicts;
    环节/是否通用取第一个非空值,后续行给出不同的非空值记 attr_diffs(审核可见,不阻断)。"""
    out: dict[str, SpuRow] = {}
    order: list[str] = []
    units: dict[str, list[str]] = collections.defaultdict(list)
    for r in rows:
        name, spec = r.name, ""
        merged = merge_map.get(hard(r.name))
        if merged is not None:
            name, spec = merged
        k = hard(name)
        s = out.get(k)
        if s is None:
            s = SpuRow(
                name=name, spu_code=mining_spu_code(name), purchase_note=r.spec_text,
                section=r.section, generic=r.generic,
            )
            out[k] = s
            order.append(k)
        else:
            for label, attr, new in (("环节", "section", r.section), ("是否通用", "generic", r.generic)):
                cur = getattr(s, attr)
                if not new:
                    continue
                if not cur:
                    setattr(s, attr, new)
                elif new != cur:
                    s.attr_diffs.append(f"{r.file_key}:{r.row_id} {label} {new} ≠ {cur}")
        if spec and spec not in s.specs:
            s.specs.append(spec)
        scene = SCENE_BY_FILE_KEY[r.file_key]
        if scene not in s.scenes:
            s.scenes.append(scene)
        s.source_ids.append(f"{r.file_key}:{r.row_id}")
        if r.unit:
            units[k].append(r.unit)
        rule = match_rule(r, rules)
        if rule is not None:
            if s.zone_cat is None:
                s.zone_cat, s.category_code = rule.zone_cat, rule.category_code
                s.rule_desc = f"{rule.file_key}/{rule.sheet}/{rule.section or '*'}/{rule.keyword or '-'}"
            elif (rule.zone_cat, rule.category_code) != (s.zone_cat, s.category_code):
                s.conflicts.append(
                    f"{r.file_key}:{r.row_id} → {rule.zone_cat}/{rule.category_code} ≠ {s.zone_cat}/{s.category_code}"
                )
    for k in order:
        us = units[k]
        out[k].unit = collections.Counter(us).most_common(1)[0][0] if us else "PCS"
    return [out[k] for k in order]


# ── DB 落库 ───────────────────────────────────────────────

from sqlalchemy import func, select  # noqa: E402
from sqlalchemy.ext.asyncio import AsyncSession  # noqa: E402

from app.db.models.category import Category  # noqa: E402
from app.db.models.product import Product, ProductStatus, ProductVisibility  # noqa: E402
from app.db.models.product_attr import ProductAttr  # noqa: E402
from app.db.models.product_sku import ProductSku, SkuStatus  # noqa: E402
from app.db.models.zone import (  # noqa: E402
    ZONE_CATEGORY_KIND_CATEGORY,
    ZONE_CATEGORY_KIND_SCENE,
    Zone,
    ZoneCategory,
    ZoneProduct,
)
from app.core.i18n_write import apply_i18n_create, apply_i18n_prefilled  # noqa: E402
from app.core.locale import SUPPORTED_LOCALES  # noqa: E402

SOURCE_LANG = "zh"


async def mark_product_i18n(product: Product) -> None:
    """脚本建的商品不会经过 API,这里走同一套 i18n 入口记账,
    让生产的定时扫描器(sweep_pending 按 i18n_pending_at)和 backfill_translations 自动补 en/sw。"""
    product.source_lang = SOURCE_LANG
    await apply_i18n_create(product, "name", product.name_zh, SOURCE_LANG)
    if product.description_zh:
        await apply_i18n_create(product, "description", product.description_zh, SOURCE_LANG)


async def mark_sku_i18n(sku: ProductSku) -> None:
    sku.source_lang = SOURCE_LANG
    await apply_i18n_create(sku, "name", sku.name_zh, SOURCE_LANG)


async def mark_attr_i18n(attr: ProductAttr, *, value_language_neutral: bool = False) -> None:
    """attr_key 的 en/sw 固定用内部键(spec/section/...),标 manual 不翻:购物车变体指纹按 attr_key_en 算,
    被机器翻成 Specification 会把同一商品的规格维度拆散。attr_value 规格值语言中性也标 manual;其余待翻。"""
    attr.source_lang = SOURCE_LANG
    key_en = attr.attr_key_en or attr.attr_key_zh
    apply_i18n_prefilled(attr, "attr_key", {"zh": attr.attr_key_zh, "en": key_en, "sw": key_en}, SOURCE_LANG)
    if value_language_neutral:
        apply_i18n_prefilled(attr, "attr_value", {loc: attr.attr_value_zh for loc in SUPPORTED_LOCALES}, SOURCE_LANG)
    else:
        await apply_i18n_create(attr, "attr_value", attr.attr_value_zh, SOURCE_LANG)


async def ensure_platform_categories(db: AsyncSession, stats: collections.Counter) -> None:
    """建 33/34/35 子树,整棵 is_active=False。code 已被别的名字占用则中止(跨环境 code 不一致防线)。"""
    codes = [c[0] for c in PLATFORM_CATEGORIES]
    existing = {c.code: c for c in (await db.execute(select(Category).where(Category.code.in_(codes)))).scalars().all()}
    l1_sort = (await db.execute(select(func.max(Category.sort_order)).where(Category.level == 1))).scalar() or 0
    for i, (code, zh, en, sw, parent) in enumerate(PLATFORM_CATEGORIES):
        cur = existing.get(code)
        if cur is not None:
            if cur.name_zh != zh:
                raise RuntimeError(f"品类 code {code} 已存在且名称不同:{cur.name_zh!r} ≠ {zh!r},请先人工处理")
            stats["platform_categories_existing"] += 1
            continue
        level = 1 if parent is None else 2
        cat = Category(
            code=code, level=level, parent_code=parent,
            sort_order=(l1_sort + 1 + i) if level == 1 else i,
            is_active=False, is_leaf=(level == 2), source_lang=SOURCE_LANG,
        )
        apply_i18n_prefilled(cat, "name", {"zh": zh, "en": en, "sw": sw}, SOURCE_LANG)
        db.add(cat)
        stats["platform_categories_created"] += 1
    await db.flush()


async def validate_category_codes(db: AsyncSession, spus: list[SpuRow]) -> None:
    codes = sorted({s.category_code for s in spus if s.category_code})
    found = {c for (c,) in (await db.execute(select(Category.code).where(Category.code.in_(codes)))).all()}
    missing = [c for c in codes if c not in found]
    if missing:
        raise RuntimeError(f"rules.csv 里的 category_code 在 categories 表不存在,FK 会失败:{missing}")


async def get_or_create_zone(db: AsyncSession, stats: collections.Counter) -> Zone:
    zone = (await db.execute(select(Zone).where(Zone.code == ZONE_CODE))).scalar_one_or_none()
    if zone is None:
        zone = Zone(code=ZONE_CODE, name_zh=ZONE_NAME[0], name_en=ZONE_NAME[1], name_sw=ZONE_NAME[2], status="ACTIVE")
        db.add(zone)
        await db.flush()
        stats["zone_created"] += 1
    return zone


async def upsert_zone_categories(db: AsyncSession, zone: Zone, stats: collections.Counter) -> dict[str, ZoneCategory]:
    out: dict[str, ZoneCategory] = {}
    specs = [(c, zh, en, sw, ZONE_CATEGORY_KIND_CATEGORY) for c, zh, en, sw in ZONE_CATEGORIES]
    specs += [(c, zh, en, sw, ZONE_CATEGORY_KIND_SCENE) for c, zh, en, sw in ZONE_SCENES]
    for i, (code, zh, en, sw, kind) in enumerate(specs):
        zc = (await db.execute(
            select(ZoneCategory).where(ZoneCategory.zone_id == zone.id, ZoneCategory.code == code)
        )).scalar_one_or_none()
        if zc is None:
            zc = ZoneCategory(zone_id=zone.id, code=code, name_zh=zh, name_en=en, name_sw=sw, kind=kind, sort_order=i)
            db.add(zc)
            stats["zone_categories_created"] += 1
        else:
            zc.name_zh, zc.name_en, zc.name_sw, zc.kind, zc.sort_order = zh, en, sw, kind, i
            stats["zone_categories_updated"] += 1
        await db.flush()
        out[code] = zc
    return out


def _spu_attrs(s: SpuRow) -> list[tuple[str, str, str]]:
    """SPU 级非可选属性 (key_zh, key_en, value)。"""
    attrs = []
    if s.section:
        attrs.append(("环节/类别", "section", s.section))
    if s.generic:
        attrs.append(("是否通用", "interchangeable", s.generic))
    return attrs


def expected_sku_codes(s: SpuRow) -> set[str]:
    """一个 SPU 按当前合并表应有的 SKU code 集合(≥2 个规格走变体,否则单默认 SKU)。"""
    if len(s.specs) >= 2:
        return {mining_sku_code(s.spu_code, spec) for spec in s.specs}
    return {mining_sku_code(s.spu_code, None)}


def _source_meta(s: SpuRow) -> dict:
    meta = {"batch": BATCH_ID, "source_ids": s.source_ids, "scenes": s.scenes}
    if s.purchase_note:
        meta["purchase_note"] = s.purchase_note   # 运营详情读这个键(operator_products.get_product)
    return meta


async def upsert_product(db: AsyncSession, s: SpuRow, stats: collections.Counter, sku_mismatches: list[str]) -> Product:
    product = (await db.execute(
        select(Product).where(Product.spu_code == s.spu_code, Product.deleted_at.is_(None))
    )).scalar_one_or_none()
    if product is not None:
        product.name_zh = s.name
        # 描述留空(含旧译文),等正式商品描述(九云/运营)再补;采购说明已移到 source_meta.purchase_note
        product.description_zh = product.description_en = product.description_sw = None
        product.category_code = s.category_code
        product.moq_unit = s.unit
        product.visibility = ProductVisibility.ZONE_ONLY
        product.status = ProductStatus.ACTIVE
        product.source_meta = _source_meta(s)
        if not product.name_en and not (product.trans_meta or {}):
            # 早期导入未打翻译标记的行:补标记一次(已有 en 或已有记账的不动,避免重复翻译)
            await mark_product_i18n(product)
            stats["products_marked_for_i18n"] += 1
        # 更新路径不动 SKU/属性。规格集合变了(合并表加/删规格)就不是"幂等重跑",必须让人看到而不是报成功。
        live = set((await db.execute(
            select(ProductSku.sku_code).where(ProductSku.product_id == product.id, ProductSku.deleted_at.is_(None))
        )).scalars().all())
        expected = expected_sku_codes(s)
        if live != expected:
            # 数量可能相等但集合不同(改了某个规格文字),所以报缺/多的 code 数而不是总数
            sku_mismatches.append(
                f"{s.spu_code} {s.name}: 库内 SKU 缺 {len(expected - live)} 个、多 {len(live - expected)} 个(当前规格 {s.specs})"
            )
        stats["products_updated"] += 1
        await db.flush()
        return product

    product = Product(
        spu_code=s.spu_code, name_zh=s.name, description_zh=None,
        category_code=s.category_code,
        status=ProductStatus.ACTIVE, visibility=ProductVisibility.ZONE_ONLY,
        unit="PCS", moq=1, moq_unit=s.unit, source=SOURCE,
        source_meta=_source_meta(s),
    )
    await mark_product_i18n(product)
    db.add(product)
    await db.flush()

    if len(s.specs) >= 2:
        stats["products_variant"] += 1
        for i, spec in enumerate(s.specs):
            sku = ProductSku(
                product_id=product.id, sku_code=mining_sku_code(s.spu_code, spec),
                name_zh=f"{s.name} {spec}", moq=1, is_default=(i == 0), status=SkuStatus.ACTIVE,
            )
            await mark_sku_i18n(sku)
            db.add(sku)
            await db.flush()
            spec_attr = ProductAttr(
                product_id=product.id, sku_id=sku.id,
                attr_key_zh="规格", attr_key_en="spec",
                attr_value_zh=spec, attr_value_en=spec, attr_value_sw=spec,
                selectable=True, sort_order=i,
            )
            await mark_attr_i18n(spec_attr, value_language_neutral=True)
            db.add(spec_attr)
        stats["skus_created"] += len(s.specs)
    else:
        stats["products_simple"] += 1
        sku = ProductSku(
            product_id=product.id, sku_code=mining_sku_code(s.spu_code, None),
            name_zh=s.name, moq=1, is_default=True, status=SkuStatus.ACTIVE,
        )
        await mark_sku_i18n(sku)
        db.add(sku)
        stats["skus_created"] += 1

    for i, (kz, ke, v) in enumerate(_spu_attrs(s)):
        attr = ProductAttr(
            product_id=product.id, sku_id=None,
            attr_key_zh=kz, attr_key_en=ke, attr_value_zh=v, selectable=False, sort_order=i,
        )
        await mark_attr_i18n(attr)
        db.add(attr)
    await db.flush()
    return product


async def upsert_zone_product(
    db: AsyncSession, zone: Zone, zc: ZoneCategory, product: Product, sort_order: int, stats: collections.Counter,
) -> None:
    zp = (await db.execute(
        select(ZoneProduct).where(
            ZoneProduct.zone_id == zone.id, ZoneProduct.spu_id == product.id, ZoneProduct.zone_category_id == zc.id,
        )
    )).scalar_one_or_none()
    if zp is not None:
        if (zp.sort_order, zp.source, zp.source_batch_id) != (sort_order, SOURCE, BATCH_ID):
            zp.sort_order, zp.source, zp.source_batch_id = sort_order, SOURCE, BATCH_ID
            stats["zone_products_updated"] += 1
            await db.flush()
        else:
            stats["zone_products_existing"] += 1
        return
    db.add(ZoneProduct(
        zone_id=zone.id, spu_id=product.id, zone_category_id=zc.id,
        sort_order=sort_order, source=SOURCE, source_batch_id=BATCH_ID,
    ))
    stats["zone_products_created"] += 1
    await db.flush()


async def prune_stale_zone_products(db: AsyncSession, zone: Zone, product: Product, expected_category_ids: set[int]) -> int:
    """规则/场景改动后重跑:删掉本批次写入、但已不在期望集合里的挂载行;手工挂的(非本批次)不动。"""
    from sqlalchemy import delete
    conds = [
        ZoneProduct.zone_id == zone.id,
        ZoneProduct.spu_id == product.id,
        ZoneProduct.source_batch_id == BATCH_ID,
    ]
    if expected_category_ids:
        conds.append(ZoneProduct.zone_category_id.not_in(expected_category_ids))
    res = await db.execute(delete(ZoneProduct).where(*conds))
    return res.rowcount or 0


@dataclass(frozen=True)
class OrphanRef:
    id: int
    spu_code: str
    name_zh: str


async def find_orphan_products(db: AsyncSession, computed_spu_codes: set[str]) -> list[OrphanRef]:
    """本批次以前导入、但这次按合并表算不出来的商品(典型:spu_merge.csv 改了品名 → spu_code 变了)。

    只取 id/spu_code/name_zh;先用带索引的 visibility + source 把候选压到专区导入商品,再做 JSON 比对,
    不让 17 万 XFS SPU 跟着全表扫。
    """
    rows = (await db.execute(
        select(Product.id, Product.spu_code, Product.name_zh).where(
            Product.deleted_at.is_(None),
            Product.visibility == ProductVisibility.ZONE_ONLY,
            Product.source == SOURCE,
            Product.source_meta["batch"].as_string() == BATCH_ID,
        ).order_by(Product.spu_code)
    )).all()
    return [OrphanRef(i, c, n) for i, c, n in rows if c not in computed_spu_codes]


async def soft_delete_orphans(db: AsyncSession, zone: Zone, orphans: list[OrphanRef]) -> int:
    """--prune-orphans:孤儿商品及其 SKU 一起软删 + 清其专区挂载。

    SKU 必须跟着软删:sku_code 由 spu_code+规格确定,留着活的 SKU,以后改回品名重导同一商品时
    会撞 uq_product_skus_sku_code_active。属性无 deleted_at,随商品不可见,不动。
    """
    from sqlalchemy import delete, update
    if not orphans:
        return 0
    ids = [o.id for o in orphans]
    now = func.now()
    await db.execute(delete(ZoneProduct).where(ZoneProduct.zone_id == zone.id, ZoneProduct.spu_id.in_(ids)))
    await db.execute(
        update(ProductSku).where(ProductSku.product_id.in_(ids), ProductSku.deleted_at.is_(None)).values(deleted_at=now)
    )
    await db.execute(update(Product).where(Product.id.in_(ids)).values(deleted_at=now))
    await db.flush()
    return len(ids)


def write_audit(spus: list[SpuRow], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["spu_code", "品名", "规格", "单位", "场景", "专区大类", "平台品类", "命中规则", "环节", "是否通用", "来源编号", "冲突", "属性差异", "采购说明"])
        for s in spus:
            w.writerow([
                s.spu_code, s.name, " / ".join(s.specs), s.unit, " / ".join(s.scenes),
                s.zone_cat or "【未命中】", s.category_code or "", s.rule_desc, s.section, s.generic,
                " / ".join(s.source_ids), " | ".join(s.conflicts), " | ".join(s.attr_diffs), s.purchase_note,
            ])


def write_dropped(dropped: list[DroppedRow], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["file_key", "sheet", "编号", "品名", "丢弃原因"])
        for d in dropped:
            w.writerow([d.file_key, d.sheet, d.row_id, d.name, d.reason])


async def run_import(db: AsyncSession, src_dir: Path, dry_run: bool, prune_orphans: bool = False) -> dict:
    stats: collections.Counter = collections.Counter()
    dropped: list[DroppedRow] = []
    rows = load_source_rows(src_dir, dropped)
    spus = build_spus(rows, load_merge_map(MERGE_CSV), load_rules(RULES_CSV))
    stats["source_rows"], stats["spus"], stats["rows_dropped"] = len(rows), len(spus), len(dropped)
    unmatched = [s for s in spus if s.zone_cat is None]
    conflicted = [s for s in spus if s.conflicts]
    stats["unmatched"], stats["conflicts"] = len(unmatched), len(conflicted)
    stats["attr_diffs"] = sum(1 for s in spus if s.attr_diffs)
    if dry_run:
        write_audit(spus, AUDIT_CSV)
        write_dropped(dropped, DROPPED_CSV)
    for d in dropped:
        print(f"丢弃行[{d.reason}]: {d.file_key}/{d.sheet}#{d.row_id} {d.name}")
    if unmatched or conflicted:
        msg = f"规则未命中 {len(unmatched)} 个、大类冲突 {len(conflicted)} 个,见 {AUDIT_CSV}"
        if not dry_run:
            raise RuntimeError("拒绝落库:" + msg)
        print("警告:" + msg)

    await ensure_platform_categories(db, stats)
    await validate_category_codes(db, spus)
    zone = await get_or_create_zone(db, stats)
    zcs = await upsert_zone_categories(db, zone, stats)

    # 孤儿守卫:上次导入过、这次算不出来的商品(合并表改名会触发)。不处理就会在专区里留下重复/过期商品。
    orphans = await find_orphan_products(db, {s.spu_code for s in spus})
    stats["orphan_products"] = len(orphans)
    if orphans:
        for p in orphans[:20]:
            print(f"孤儿商品: {p.spu_code} {p.name_zh}")
        if prune_orphans:
            stats["orphans_soft_deleted"] = await soft_delete_orphans(db, zone, orphans)
        elif not dry_run:
            raise RuntimeError(f"拒绝落库:{len(orphans)} 个孤儿商品(上批导入、本次不在清单内)。确认后加 --prune-orphans 软删,或改回品名。")

    sort_counters: dict[str, int] = collections.defaultdict(int)
    sku_mismatches: list[str] = []
    for s in spus:
        if s.zone_cat is None:
            continue
        product = await upsert_product(db, s, stats, sku_mismatches)
        expected: set[int] = set()
        await upsert_zone_product(db, zone, zcs[s.zone_cat], product, sort_counters[s.zone_cat], stats)
        expected.add(zcs[s.zone_cat].id)
        sort_counters[s.zone_cat] += 1
        for scene in s.scenes:
            await upsert_zone_product(db, zone, zcs[scene], product, sort_counters[scene], stats)
            expected.add(zcs[scene].id)
            sort_counters[scene] += 1
        # 规则改动后重跑:清掉本批次写入但已不在期望集合里的挂载
        stats["zone_products_removed"] += await prune_stale_zone_products(db, zone, product, expected)

    stats["sku_mismatches"] = len(sku_mismatches)
    if sku_mismatches:
        for m in sku_mismatches[:20]:
            print(f"SKU 不一致: {m}")
        msg = f"{len(sku_mismatches)} 个已存在商品的规格集合变了,脚本不自动增删 SKU(会动已有引用),需人工处理"
        if not dry_run:
            raise RuntimeError("拒绝落库:" + msg)
        print("警告:" + msg)
    return dict(stats)


async def _execute(src_dir: Path, dry_run: bool, prune_orphans: bool = False) -> None:
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
    from app.core.config import settings

    guard_commit(settings.DATABASE_URL, dry_run=dry_run)

    engine = create_async_engine(settings.DATABASE_URL)
    Session = async_sessionmaker(engine, expire_on_commit=False)
    async with Session() as db:
        stats = await run_import(db, src_dir, dry_run, prune_orphans=prune_orphans)
        if dry_run:
            await db.rollback()
        else:
            await db.commit()
    await engine.dispose()

    print("=" * 64)
    print(f"矿业专区导入 {'[DRY-RUN 未落库]' if dry_run else '[已落库 COMMIT]'}  batch={BATCH_ID}")
    print("=" * 64)
    for k in ("source_rows", "rows_dropped", "spus", "unmatched", "conflicts", "attr_diffs", "sku_mismatches",
              "platform_categories_created", "platform_categories_existing",
              "zone_created", "zone_categories_created", "zone_categories_updated",
              "products_simple", "products_variant", "products_updated", "products_marked_for_i18n", "skus_created",
              "zone_products_created", "zone_products_updated", "zone_products_existing", "zone_products_removed",
              "orphan_products", "orphans_soft_deleted"):
        print(f"{k:32s} {stats.get(k, 0)}")
    if dry_run:
        print(f"\n审核 CSV:{AUDIT_CSV}\n丢弃行 CSV:{DROPPED_CSV}\n这是干跑,数据库未改动。未命中/冲突/SKU 不一致清零后加 --commit 落库。")


def main() -> None:
    ap = argparse.ArgumentParser(description="矿业勘探·开采专区 6 张配货表导入")
    ap.add_argument("--src", type=Path, required=True, help="6 个 xlsx 所在目录")
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--dry-run", action="store_true")
    g.add_argument("--commit", action="store_true")
    ap.add_argument("--prune-orphans", action="store_true", help="把上批导入、本次不在清单内的商品软删并清其专区挂载(合并表改名后用)")
    args = ap.parse_args()
    asyncio.run(_execute(args.src, dry_run=args.dry_run, prune_orphans=args.prune_orphans))


if __name__ == "__main__":
    main()
