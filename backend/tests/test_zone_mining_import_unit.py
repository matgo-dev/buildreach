"""矿业专区导入脚本纯函数单测:表格解析 / 品名合并 / 规则匹配 / 编码稳定性。不碰 DB。"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.services.product_code import platform_spu_code
from scripts.import_zone_mining import (
    Rule,
    SourceRow,
    build_spus,
    hard,
    load_merge_map,
    load_rules,
    load_source_rows,
    match_rule,
    mining_spu_code,
)


# 单测用的平台品类 → 专区大类对照(与 zone_category_map.csv 同形)
CMAP = {"01": "C10", "04": "C09", "14": "C06", "33.001": "C01", "34": "C04", "35": "C11"}


def _write_book(path: Path, sheets: dict[str, list[list]]) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    wb.save(path)


def test_hard_normalizes_fullwidth_and_whitespace():
    assert hard("８φ 绳卡　") == "8φ绳卡"


def test_mining_spu_code_is_platform_neutral_and_stable():
    code = mining_spu_code("绳卡")
    assert code == platform_spu_code("ZMINE", "绳卡")
    assert code.startswith("MG-P") and len(code) == 16
    assert "ZMINE" not in code and "绳卡" not in code


def test_load_source_rows_handles_both_header_variants(tmp_path: Path):
    _write_book(tmp_path / "01_浮选厂.xlsx", {
        "浮选设备与备件": [
            ["标题行"], [],
            ["编号", "环节/类别", "设备、配件及易耗品", "参考型号/规格", "单位", "100吨", "500吨", "a", "b", "c", "d", "是否通用"],
            ["F-01", "备件", "颚式破碎机颚板", "按颚破型号、图号及齿形", "套", 1, 1, None, None, None, None, "否（机型专用）"],
            [None, None, None, None, None],
        ],
        "规模与汇总": [["不导入"]],
    })
    _write_book(tmp_path / "600型与800型.xlsx", {
        "600型10台预算": [
            ["编号", "配件名称", "型号或规格", "单位", "半年参考量"],
            ["新R007", "钢丝绳", "6mm×1000m", "捆", None],
        ],
        "800型10台预算": [["编号", "配件名称"], ["新R007", "钢丝绳"]],
    })
    rows = load_source_rows(tmp_path)
    assert [(r.file_key, r.sheet, r.row_id, r.name) for r in rows] == [
        ("01", "浮选设备与备件", "F-01", "颚式破碎机颚板"),
        ("600", "600型10台预算", "新R007", "钢丝绳"),
    ]
    assert rows[0].section == "备件"
    assert rows[0].spec_text == "按颚破型号、图号及齿形"
    assert rows[0].unit == "套"
    assert rows[0].generic == "否（机型专用）"
    assert rows[1].section == ""  # 600 表没有类别列


def test_load_merge_map_and_blank_spec_is_alias(tmp_path: Path):
    p = tmp_path / "spu_merge.csv"
    p.write_text("原品名,SPU品名,规格\n8φ绳卡,绳卡,φ8\n颚破边护板,颚式破碎机边护板,\n", encoding="utf-8")
    m = load_merge_map(p)
    assert m[hard("8φ绳卡")] == ("绳卡", "φ8")
    assert m[hard("颚破边护板")] == ("颚式破碎机边护板", "")


def test_match_rule_keyword_rules_win_over_sheet_defaults():
    rules = [
        Rule("01", "药剂及钢球", "*", "", "34.001"),
        Rule("01", "药剂及钢球", "*", "钢球", "33.001"),
        Rule("01", "*", "劳保", "", "01"),
    ]
    ball = SourceRow("01", "药剂及钢球", "X-1", "", "研磨钢球", "", "吨", "")
    agent = SourceRow("01", "药剂及钢球", "X-2", "", "捕收剂", "", "kg", "")
    glove = SourceRow("01", "浮选现场配套", "X-3", "劳保", "帆布手套", "", "双", "")
    other = SourceRow("02", "井下设备备件", "X-4", "", "钎杆", "", "根", "")
    assert match_rule(ball, rules).category_code == "33.001"
    assert match_rule(agent, rules).category_code == "34.001"
    assert match_rule(glove, rules).category_code == "01"
    assert match_rule(other, rules) is None


def test_load_rules_reads_csv(tmp_path: Path):
    p = tmp_path / "rules.csv"
    p.write_text("file_key,sheet,section,keyword,category_code,note\n01,*,*,,01,x\n", encoding="utf-8")
    rules = load_rules(p)
    assert rules == [Rule("01", "*", "*", "", "01")]


def _rules_default():
    return [
        Rule("01", "*", "*", "", "01"),
        Rule("02", "*", "*", "", "04"),
    ]


def test_build_spus_merges_sizes_into_skus_and_unions_scenes():
    rows = [
        SourceRow("01", "浮选现场配套", "A-1", "耗材", "8φ绳卡", "原文φ8", "个", "限同制式或接口"),
        SourceRow("01", "浮选现场配套", "A-2", "耗材", "15φ绳卡", "原文φ15", "个", "限同制式或接口"),
        SourceRow("02", "井下设备备件", "B-1", "耗材", "8φ绳卡", "原文φ8", "个", "限同制式或接口"),
    ]
    merge = {hard("8φ绳卡"): ("绳卡", "φ8"), hard("15φ绳卡"): ("绳卡", "φ15")}
    rules = [Rule("*", "*", "*", "", "04")]
    spus = build_spus(rows, merge, rules, CMAP)
    assert len(spus) == 1
    s = spus[0]
    assert s.name == "绳卡"
    assert s.spu_code == mining_spu_code("绳卡")
    assert s.specs == ["φ8", "φ15"]
    assert s.scenes == ["S01", "S02"]
    assert s.purchase_note == "原文φ8"
    assert s.source_ids == ["01:A-1", "01:A-2", "02:B-1"]
    assert s.zone_cat == "C09" and s.category_code == "04"
    assert s.conflicts == []


def test_build_spus_reports_category_conflict():
    rows = [
        SourceRow("01", "浮选现场配套", "A-1", "劳保", "帆布手套", "", "双", ""),
        SourceRow("02", "全矿共用维修工具", "B-9", "工具", "帆布手套", "", "双", ""),
    ]
    spus = build_spus(rows, {}, _rules_default(), CMAP)
    assert len(spus) == 1
    assert spus[0].category_code == "01" and spus[0].zone_cat == "C10"      # 第一次命中为准,大类由品类推导
    assert spus[0].conflicts == ["02:B-9 → 04 ≠ 01"]


def test_build_spus_unmatched_row_has_no_zone_cat():
    rows = [SourceRow("03", "两法共用", "LAB-1", "器皿", "烧杯", "", "个", "")]
    spus = build_spus(rows, {}, _rules_default(), CMAP)
    assert spus[0].zone_cat is None and spus[0].category_code is None


def test_load_source_rows_skips_non_product_noise_rows(tmp_path: Path):
    """「待核数量」「xx方案」「xx服务」这类占位/服务行不是商品,解析时丢弃。"""
    _write_book(tmp_path / "05_露天.xlsx", {
        "油品及运行耗材": [
            ["编号", "类别", "设备、配件及易耗品", "参考型号/规格", "单位"],
            ["OP-O001", "油品", "柴油机油", "按厂家规格", "桶"],
            ["OP-X001", "3台挖机方案", "5台挖机方案", "", ""],
            ["OP-X002", "待核数量", "待核数量", "", ""],
            ["OP-X003", "救援", "车辆拖救服务", "", "次"],
        ],
    })
    rows = load_source_rows(tmp_path)
    assert [r.name for r in rows] == ["柴油机油"]


def test_match_rule_file_specific_keyword_beats_wildcard_keyword():
    """同为关键词规则时,写了 file_key 的比 * 的优先,不管谁在前(化验室的「pH计」不能被全局「仪/计」抢走)。"""
    rules = [
        Rule("*", "*", "*", "计", "14"),
        Rule("03", "*", "*", "计", "35.001"),
    ]
    lab = SourceRow("03", "两法共用", "LAB-1", "仪器", "实验室pH计", "", "台", "")
    site = SourceRow("04", "回收及公辅", "R-1", "", "流量计", "", "台", "")
    assert match_rule(lab, rules).category_code == "35.001"
    assert match_rule(site, rules).category_code == "14"


@pytest.mark.asyncio
async def test_mark_i18n_sets_pending_for_translatable_fields_and_manual_for_spec():
    """导入行必须打翻译标记,否则生产的定时扫描器(sweep_pending 按 i18n_pending_at)永远不会补 en/sw。"""
    from app.db.models.product import Product
    from app.db.models.product_attr import ProductAttr
    from scripts.import_zone_mining import mark_attr_i18n, mark_product_i18n

    p = Product(spu_code="X", name_zh="绳卡", description_zh="按钢丝绳直径核对", category_code="04.001")
    await mark_product_i18n(p)
    assert p.source_lang == "zh"
    assert p.trans_meta["name_zh"] == "src" and p.trans_meta["name_en"] == "pending" and p.trans_meta["name_sw"] == "pending"
    assert p.trans_meta["description_en"] == "pending"
    assert p.i18n_pending_at is not None

    spec = ProductAttr(attr_key_zh="规格", attr_value_zh="扬程150米", selectable=True)
    await mark_attr_i18n(spec)
    # 可选规格的 en 键是交易键(购物车指纹按 attr_key_en 算):固定、不机翻;sw 只展示,走翻译
    assert spec.trans_meta["attr_key_en"] == "manual" and spec.attr_key_en == "Specification"
    assert spec.trans_meta["attr_key_sw"] == "pending"
    # 规格值走翻译(「扬程150米」这类中文不能原样给英文买家)
    assert spec.trans_meta["attr_value_en"] == "pending" and spec.attr_value_en is None and spec.attr_value_sw is None

    plain = ProductAttr(attr_key_zh="是否通用", attr_value_zh="否（机型专用）", selectable=False)
    await mark_attr_i18n(plain)
    # 非可选属性的键是展示名,按中文走翻译(买家不能看到内部键)
    assert plain.trans_meta["attr_key_en"] == "pending" and plain.attr_key_en is None
    assert plain.trans_meta["attr_value_en"] == "pending"
    assert plain.i18n_pending_at is not None


def test_noise_filter_drops_report_forms():
    from scripts.import_zone_mining import is_noise_name
    assert is_noise_name("班报表")
    assert not is_noise_name("岩心牌")


def test_build_spus_records_matched_rule_for_audit():
    rows = [SourceRow("01", "药剂及钢球", "X-1", "", "研磨钢球", "", "吨", "")]
    rules = [Rule("*", "*", "*", "钢球", "33.001")]
    assert build_spus(rows, {}, rules, CMAP)[0].rule_desc == "*/*/*/钢球"


def test_load_source_rows_records_dropped_rows_with_reason(tmp_path: Path):
    """噪声/超长行不能静默消失:计数 + 原因,dry-run 写进丢弃 CSV 给人复核。"""
    from scripts.import_zone_mining import DroppedRow
    long_name = "高压油管总成" * 12
    _write_book(tmp_path / "01_浮选厂.xlsx", {"药剂": [
        ["编号", "易耗品名称", "单位"],
        ["1", "研磨钢球", "吨"],
        ["2", "车辆拖救服务", "次"],
        ["3", long_name, "套"],
        ["", "", ""],
    ]})
    dropped: list[DroppedRow] = []
    rows = load_source_rows(tmp_path, dropped)
    assert [r.name for r in rows] == ["研磨钢球"]
    assert [(d.row_id, d.reason) for d in dropped] == [("2", "noise"), ("3", "too_long")]
    # 不传 dropped 也照常工作
    assert len(load_source_rows(tmp_path)) == 1


def test_build_spus_records_attr_diffs_across_files_without_blocking():
    """同一 SPU 在两个文件里环节不同:记 attr_diffs 供审核,取第一个非空值,不算大类冲突。"""
    rules = [Rule("*", "*", "*", "手套", "01")]
    rows = [
        SourceRow("01", "劳保", "A-1", "劳保", "帆布手套", "", "双", "是"),
        SourceRow("02", "工具", "B-7", "工具", "帆布手套", "", "双", ""),
        SourceRow("05", "耗材", "E-2", "", "帆布手套", "", "双", "限同制式"),
    ]
    (s,) = build_spus(rows, {}, rules, CMAP)
    assert s.section == "劳保" and s.generic == "是"
    assert s.attr_diffs == ["02:B-7 环节 工具 ≠ 劳保", "05:E-2 是否通用 限同制式 ≠ 是"]
    assert s.conflicts == []


def test_build_spus_fills_empty_attr_from_later_row():
    rules = [Rule("*", "*", "*", "手套", "01")]
    rows = [
        SourceRow("01", "劳保", "A-1", "", "帆布手套", "", "双", ""),
        SourceRow("02", "工具", "B-7", "工具", "帆布手套", "", "双", "是"),
    ]
    (s,) = build_spus(rows, {}, rules, CMAP)
    assert (s.section, s.generic, s.attr_diffs) == ("工具", "是", [])


def test_expected_sku_codes_variant_vs_simple():
    from scripts.import_zone_mining import SpuRow, expected_sku_codes, mining_sku_code
    simple = SpuRow(name="活板", spu_code="MG-PX", purchase_note="", specs=["20寸"])
    assert expected_sku_codes(simple) == {mining_sku_code("MG-PX", None)}
    variant = SpuRow(name="活板", spu_code="MG-PX", purchase_note="", specs=["20寸", "24寸"])
    assert expected_sku_codes(variant) == {mining_sku_code("MG-PX", "20寸"), mining_sku_code("MG-PX", "24寸")}


@pytest.mark.parametrize("url,needs", [
    ("postgresql+asyncpg://u:p@localhost:5433/x", False),
    ("postgresql+asyncpg://u:p@127.0.0.1/x", False),
    ("postgresql+asyncpg://u:p@db:5432/x", True),            # OVH/ECS 容器里的 compose 服务名
    ("postgresql+asyncpg://u:p@162.19.98.142/x", True),
    ("", True),
])
def test_commit_guard_requires_opt_in_for_any_non_local_db(url, needs, monkeypatch):
    from scripts._zone_import_common import OPT_IN_ENV, commit_needs_opt_in, guard_commit
    assert commit_needs_opt_in(url) is needs
    monkeypatch.delenv(OPT_IN_ENV, raising=False)
    guard_commit(url, dry_run=True)   # dry-run 永远放行
    if needs:
        with pytest.raises(SystemExit):
            guard_commit(url, dry_run=False)
        monkeypatch.setenv(OPT_IN_ENV, "1")
    guard_commit(url, dry_run=False)


def test_normalize_source_key_shared_by_both_zone_scripts():
    """两个专区脚本的 hash 输入必须走同一份归一化,否则同名商品在两边算出不同 code。"""
    from app.services.product_code import normalize_source_key
    from scripts.import_zone_soe_materials import _hard
    for raw in ("ＡＢＣ 12 3", " 螺纹钢\t", None, 42):
        assert hard(raw) == _hard(raw) == normalize_source_key(raw)
    assert normalize_source_key(None) == ""


def test_clean_spec_text_drops_placeholder():
    """源表规格列的「未列」是"没填",不能当商品描述展示;NTW/HTW 是真实钻具口径,保留。"""
    from scripts.import_zone_mining import clean_spec_text
    assert clean_spec_text("未列") == "" and clean_spec_text(" 未列 ") == ""
    assert clean_spec_text("原PDF未列规格") == ""
    assert clean_spec_text("NTW") == "NTW" and clean_spec_text("按钻机型号") == "按钻机型号"


def test_clean_generic_drops_placeholder_and_spu_gets_no_attr():
    from scripts.import_zone_mining import _spu_attrs, clean_generic
    assert clean_generic("待确认") == "" and clean_generic("是") == "是"
    rows = [SourceRow("02", "井下", "D-1", "耗材", "起爆器", "", "个", clean_generic("待确认"))]
    (s,) = build_spus(rows, {}, [Rule("*", "*", "*", "起爆器", "01")], CMAP)
    assert [k for k, _ in _spu_attrs(s)] == ["环节/类别"]


def test_spu_attrs_never_emit_applicable_model():
    """「适用机型」曾从「按XX型号核对」里抽指令片段(颚破型号/确定机型),不是型号,已整体去掉。"""
    from scripts.import_zone_mining import _spu_attrs
    rows = [SourceRow("01", "破碎", "A-1", "备件", "颚板", "按颚破型号核对", "块", "否")]
    (s,) = build_spus(rows, {}, [Rule("*", "*", "*", "颚板", "33.001")], CMAP)
    assert "适用机型" not in [k for k, _ in _spu_attrs(s)]


def test_purchase_note_goes_to_source_meta_not_description():
    """采购说明是写给采购员的(按…核对/预算示例/待核),只进运营可见的 source_meta.purchase_note。"""
    from scripts.import_zone_mining import _source_meta
    rows = [SourceRow("01", "破碎", "A-1", "备件", "颚板", "预算示例：按颚破型号核对", "块", "否")]
    (s,) = build_spus(rows, {}, [Rule("*", "*", "*", "颚板", "33.001")], CMAP)
    assert s.purchase_note == "预算示例：按颚破型号核对"
    assert _source_meta(s)["purchase_note"] == "预算示例：按颚破型号核对"
    rows[0] = SourceRow("01", "破碎", "A-1", "备件", "颚板", "", "块", "否")
    (s,) = build_spus(rows, {}, [Rule("*", "*", "*", "颚板", "33.001")], CMAP)
    assert "purchase_note" not in _source_meta(s)


@pytest.mark.parametrize("code,expected", [
    ("33.001", "C01"), ("33.001.002", "C01"),   # 叶子继承祖先前缀
    ("34.002", "C04"), ("34", "C04"),
    ("07.002.002.001", "C06"),                  # 长前缀 07.002 优先于 07
    ("07.009", "C07"), ("07", "C07"),           # 其余落到 07
    ("33.00", None),                            # 按段匹配,不按字符:'33.00' 不是 '33.001' 的祖先
    ("99", None), (None, None), ("", None),
])
def test_derive_zone_cat_longest_segment_prefix(code, expected):
    from scripts.import_zone_mining import derive_zone_cat
    cmap = {"33.001": "C01", "34": "C04", "07": "C07", "07.002": "C06"}
    assert derive_zone_cat(code, cmap) == expected


def test_load_zone_cat_map_rejects_duplicate_prefix_and_unknown_zone_cat(tmp_path: Path):
    from scripts.import_zone_mining import load_zone_cat_map
    p = tmp_path / "m.csv"
    p.write_text("platform_code_prefix,zone_cat,note\n07,C07,x\n07.002,C06,y\n", encoding="utf-8")
    assert load_zone_cat_map(p) == {"07": "C07", "07.002": "C06"}
    p.write_text("platform_code_prefix,zone_cat,note\n07,C07,x\n07,C06,y\n", encoding="utf-8")
    with pytest.raises(ValueError, match="前缀重复"):
        load_zone_cat_map(p)
    p.write_text("platform_code_prefix,zone_cat,note\n07,C99,x\n", encoding="utf-8")
    with pytest.raises(ValueError, match="未知专区大类"):
        load_zone_cat_map(p)


def test_build_spus_category_without_mapping_leaves_zone_cat_empty():
    """规则给了平台品类、对照表没覆盖 → zone_cat=None(run_import 计 unmapped_zone_cat 并拒绝落库),不能悄悄归某个大类。"""
    rows = [SourceRow("01", "x", "A-1", "", "某商品", "", "个", "")]
    (s,) = build_spus(rows, {}, [Rule("*", "*", "*", "", "99")], CMAP)
    assert s.category_code == "99" and s.zone_cat is None


def test_real_zone_category_map_covers_every_rule_category():
    """真实规则表里每个平台品类都必须能推出专区大类(否则导入时整批拒绝)。"""
    from scripts.import_zone_mining import RULES_CSV, ZONE_CAT_MAP_CSV, derive_zone_cat, load_zone_cat_map
    cmap = load_zone_cat_map(ZONE_CAT_MAP_CSV)
    missing = sorted({r.category_code for r in load_rules(RULES_CSV) if derive_zone_cat(r.category_code, cmap) is None})
    assert missing == []


def test_load_merge_map_rejects_duplicate_source_name(tmp_path: Path):
    p = tmp_path / "m.csv"
    p.write_text("原品名,SPU品名,规格\n8φ绳卡,绳卡,φ8\n８φ绳卡,绳卡,φ18\n", encoding="utf-8")   # 全角数字归一后同名
    with pytest.raises(ValueError, match="重复"):
        load_merge_map(p)


def test_build_spus_blocks_single_spec_that_would_be_dropped():
    rows = [SourceRow("01", "浮选现场配套", "A-1", "耗材", "20φ钢丝绳", "", "米", "")]
    spus = build_spus(rows, {hard("20φ钢丝绳"): ("钢丝绳", "Φ20")}, [Rule("*", "*", "*", "", "04")], CMAP)
    assert spus[0].conflicts and "只有一个规格" in spus[0].conflicts[0]


def test_load_source_rows_skips_unknown_file_keys(tmp_path: Path):
    import openpyxl
    from scripts.import_zone_mining import load_source_rows
    wb = openpyxl.Workbook()
    wb.active.append(["编号", "品名"])
    wb.active.append(["X-1", "绳卡"])
    wb.save(tmp_path / "06_汇总.xlsx")
    assert load_source_rows(tmp_path) == []


def test_spec_value_was_prefilled_only_matches_old_neutral_rows():
    """重跑只把「抄中文标 manual」的旧规格值改回待翻;运营手改过的英文、非可选属性都不动。"""
    from app.db.models.product_attr import ProductAttr
    from scripts.import_zone_mining import _spec_value_was_prefilled

    old = ProductAttr(attr_key_zh="规格", attr_value_zh="扬程150米", attr_value_en="扬程150米", selectable=True,
                      trans_meta={"attr_value_en": "manual"})
    edited = ProductAttr(attr_key_zh="规格", attr_value_zh="扬程150米", attr_value_en="Head 150 m", selectable=True,
                         trans_meta={"attr_value_en": "manual"})
    pending = ProductAttr(attr_key_zh="规格", attr_value_zh="扬程150米", selectable=True, trans_meta={"attr_value_en": "pending"})
    plain = ProductAttr(attr_key_zh="是否通用", attr_value_zh="否", attr_value_en="否", selectable=False,
                        trans_meta={"attr_value_en": "manual"})
    # 更早的本地导入连 attr_value 标记都没打,只抄了 en
    unmarked = ProductAttr(attr_key_zh="规格", attr_value_zh="扬程150米", attr_value_en="扬程150米", selectable=True,
                           trans_meta={"attr_key_en": "manual"})
    # 翻译器处理过的中性值(译完仍等于中文)不能反复送译
    neutral_done = ProductAttr(attr_key_zh="规格", attr_value_zh="NTW", attr_value_en="NTW", selectable=True,
                               trans_meta={"attr_value_en": "auto"})
    assert _spec_value_was_prefilled(old)
    assert _spec_value_was_prefilled(unmarked)
    assert not _spec_value_was_prefilled(edited)
    assert not _spec_value_was_prefilled(pending)
    assert not _spec_value_was_prefilled(plain)
    assert not _spec_value_was_prefilled(neutral_done)
